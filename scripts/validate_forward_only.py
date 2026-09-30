"""Fail closed when the ratified CHANGE-067 boundary is incomplete or mutated."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from io import BytesIO
from pathlib import Path
from typing import Any, Mapping

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from scripts.activate_forward_only import ActivationSourceError, _ledger_rows, _read_json  # noqa: E402
from src import forward_only, forward_only_projections  # noqa: E402


def _xlsx_cells(value: bytes) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(BytesIO(value), data_only=False, read_only=False)
    return {
        sheet.title: {
            cell.coordinate: cell.value
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
        for sheet in workbook.worksheets
    }


def _projection_errors(boundary: forward_only.Boundary, *, root: Path) -> list[str]:
    errors: list[str] = []
    references = (
        ((boundary.manifest or {}).get("protected") or {}).get(
            "historic_aggregate_references"
        )
        or []
    )
    for reference in references:
        if not isinstance(reference, Mapping):
            errors.append("projection:ACTIVATION_PROJECTION_INVENTORY_INVALID")
            continue
        relative = str(reference.get("path") or "")
        try:
            baseline = forward_only_projections.baseline_bytes(
                boundary, relative, root=root,
            )
        except Exception as exc:
            errors.append(f"projection:{relative}:{exc}")
            continue
        path = root / relative
        if not path.is_file():
            errors.append(f"projection:{relative}:PROTECTED_PROJECTION_MISSING")
            continue
        current = path.read_bytes()
        if str(reference.get("kind") or "").upper() == "XLSX":
            try:
                baseline_cells = _xlsx_cells(baseline)
                current_cells = _xlsx_cells(current)
            except Exception:
                errors.append(f"projection:{relative}:PROTECTED_XLSX_INVALID")
                continue
            for sheet, cells in baseline_cells.items():
                if sheet not in current_cells or current_cells[sheet] != cells:
                    errors.append(
                        f"projection:{relative}:PROTECTED_XLSX_HISTORIC_COMPONENT_MUTATED"
                    )
                    break
        else:
            try:
                baseline_doc = json.loads(baseline)
                current_doc = json.loads(current)
                forward_only_projections.validate_historic_component(
                    current_doc, baseline_doc, relative_path=relative,
                )
            except Exception as exc:
                errors.append(f"projection:{relative}:{exc}")
    return errors


def validate(manifest_path: Path, *, root: Path | None = None) -> list[str]:
    root = Path(root or ROOT)
    boundary = forward_only.load_boundary(manifest_path)
    if boundary.fail_closed:
        return [boundary.reason_code]
    if not boundary.active:
        return []
    manifest = boundary.manifest or {}
    protected = manifest.get("protected") or {}
    errors: list[str] = []
    try:
        current = {
            "snapshots": _read_json(root / "data/calibration_snapshots.json", "snapshots"),
            "paper": _read_json(root / "data/paper_trades.json", "entries"),
            "market_ledger": list(_ledger_rows(root / "data/market_ledger")),
        }
    except ActivationSourceError as exc:
        return [f"ACTIVATION_SOURCE_UNREADABLE:{exc}"]
    for collection, rows in current.items():
        actual = {
            forward_only.record_key(collection, row): forward_only.canonical_sha256(row)
            for row in rows
        }
        section = protected.get(collection) or {}
        for expected in section.get("records") or []:
            identity = str(expected.get("identity") or "")
            if identity not in actual:
                errors.append(f"{collection}:{identity}:PROTECTED_RECORD_MISSING")
            elif actual[identity] != expected.get("sha256"):
                errors.append(f"{collection}:{identity}:PROTECTED_RECORD_MUTATED")
    for report in (protected.get("published_reports") or {}).get("files") or []:
        path = root / str(report.get("path") or "")
        if not path.exists():
            errors.append(f"report:{report.get('path')}:PROTECTED_REPORT_MISSING")
        elif hashlib.sha256(path.read_bytes()).hexdigest() != report.get("sha256"):
            errors.append(f"report:{report.get('path')}:PROTECTED_REPORT_MUTATED")
    for section in protected.get("exclusions") or []:
        relative = str(section.get("path") or "")
        try:
            rows = _read_json(root / relative, "exclusions")
        except ActivationSourceError as exc:
            errors.append(f"exclusion:{relative}:PROTECTED_SOURCE_UNREADABLE:{exc}")
            continue
        actual = {}
        for index, row in enumerate(rows):
            identity = str(
                row.get("paper_key") or row.get("snapshot_key")
                or row.get("report_id") or f"index:{index}"
            )
            actual[identity] = forward_only.canonical_sha256(row)
        for expected in section.get("records") or []:
            identity = str(expected.get("identity") or "")
            if identity not in actual:
                errors.append(f"exclusion:{identity}:PROTECTED_RECORD_MISSING")
            elif actual[identity] != expected.get("sha256"):
                errors.append(f"exclusion:{identity}:PROTECTED_RECORD_MUTATED")
    errors.extend(_projection_errors(boundary, root=root))
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, default=forward_only.DEFAULT_MANIFEST_PATH)
    args = parser.parse_args()
    errors = validate(args.manifest)
    if errors:
        print(json.dumps({"status": "FAILED", "reason_codes": errors[:100]}, indent=2))
        return 1
    boundary = forward_only.load_boundary(args.manifest)
    print(json.dumps({
        "status": "VALIDATED" if boundary.active else "EXTERNAL_ACTIVATION_PENDING",
        "manifest": str(args.manifest),
    }))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
