"""Build the downloadable, canonical Fenzobot system-history workbook.

This command has no network clients.  It only reads artefacts already
persisted by the bot and leaves the existing workbook untouched if their
semantic fingerprint has not changed.
"""

from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.system_history_analytics import build_system_history, load_local_wta_history, write_if_changed
from src import forward_only, forward_only_projections


SNAPSHOTS = ROOT / "data" / "calibration_snapshots.json"
HISTORY_CACHE = ROOT / "data" / "history_cache"
REPORTS_DIR = ROOT / "docs" / "relatorios"
ANALYTICS_JSON = ROOT / "data" / "dashboard" / "system_history_analytics.json"
WORKBOOK_PATH = ROOT / "docs" / "dashboard" / "Fenzobot_Historico_do_Sistema.xlsx"

NAVY = "17365D"
TEAL = "0F766E"
LIGHT_BLUE = "EAF2F8"
PALE_GREEN = "E8F5E9"
PALE_AMBER = "FFF4D6"
PALE_RED = "FDECEC"


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, json.JSONDecodeError):
        return {"snapshots": []}
    return value if isinstance(value, dict) else {"snapshots": []}


def _header(ws, row: int, labels: Iterable[str]) -> None:
    values = list(labels)
    for column, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def _title(ws, title: str, subtitle: str) -> None:
    ws["A2"] = title
    ws["A2"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
    ws["A3"] = subtitle
    ws["A3"].font = Font(name="Arial", size=10, italic=True, color="5B6573")


def _write_rows(ws, start_row: int, headers: list[str], rows: Iterable[Mapping[str, Any]], fields: list[str]) -> int:
    _header(ws, start_row, headers)
    current = start_row + 1
    for item in rows:
        for column, field in enumerate(fields, 1):
            ws.cell(current, column, item.get(field))
        current += 1
    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{max(start_row, current - 1)}"
    for row in ws.iter_rows(min_row=start_row, max_row=current - 1, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return current - 1


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.sheet_view.showGridLines = False


def _percent_columns(ws, row_start: int, row_end: int, columns: Iterable[int]) -> None:
    for column in columns:
        for row in range(row_start, row_end + 1):
            value = ws.cell(row, column).value
            if isinstance(value, (int, float)):
                ws.cell(row, column).value = value / 100
                ws.cell(row, column).number_format = "0.0%"


def _write_rankings(ws, rankings: Mapping[str, Mapping[str, Any]]) -> None:
    """Compact Top/Bottom 10 blocks; the sample threshold lives in the JSON."""
    _title(ws, "Rankings de aprendizagem", "Top 10 e Bottom 10 por métrica. Só entram linhas com a amostra mínima indicada; são leituras factuais, não sinais.")
    row = 5
    for category, blocks in rankings.items():
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row, 1, f"{category} · amostra mínima {blocks.get('minimum_sample', '—')}")
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        row += 1
        _header(ws, row, ["Top 10", "Amostra", "Métrica", "", "Bottom 10", "Amostra", "Métrica", ""])
        strongest = blocks.get("strongest", [])
        weakest = blocks.get("weakest", [])
        length = max(len(strongest), len(weakest), 1)
        for offset in range(length):
            current = row + 1 + offset
            if offset < len(strongest):
                item = strongest[offset]
                ws.cell(current, 1, item["label"])
                ws.cell(current, 2, item["sample"])
                ws.cell(current, 3, item["metric_pct"] / 100)
            elif offset == 0:
                ws.cell(current, 1, "Sem amostras suficientes")
            if offset < len(weakest):
                item = weakest[offset]
                ws.cell(current, 5, item["label"])
                ws.cell(current, 6, item["sample"])
                ws.cell(current, 7, item["metric_pct"] / 100)
            elif offset == 0:
                ws.cell(current, 5, "Sem amostras suficientes")
            for column in (3, 7):
                ws.cell(current, column).number_format = "0.0%"
            ws.cell(current, 1).fill = PatternFill("solid", fgColor=PALE_GREEN)
            ws.cell(current, 5).fill = PatternFill("solid", fgColor=PALE_RED)
        row += length + 3
    ws.freeze_panes = "A6"
    _widths(ws, [40, 12, 14, 4, 40, 12, 14, 4])


def _workbook(payload: Mapping[str, Any], raw_wta: list[Mapping[str, Any]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    _title(ws, "Fenzobot — Histórico Canónico", "Atualização automática baseada em snapshots imutáveis. Relatórios HTML repetidos não entram nas métricas.")
    summary = payload["summary"]
    _header(ws, 5, ["Indicador", "Valor"])
    items = [
        ("Snapshots guardados (brutos)", summary["raw_snapshots"]),
        ("Partidas canónicas", summary["canonical_snapshots"]),
        ("Repetições excluídas das métricas", summary["duplicate_snapshots_excluded"]),
        ("Versões HTML guardadas", summary["raw_report_html_versions"]),
        ("Partidas canónicas liquidadas", summary["settled_canonical_snapshots"]),
        ("Jogos WTA históricos locais", summary["historical_wta_matches"]),
        ("Regra de contagem", "primeiro snapshot pré-jogo válido por partida"),
    ]
    for row, item in enumerate(items, 6):
        ws.cell(row, 1, item[0]).font = Font(name="Arial", bold=True, color="243B53")
        ws.cell(row, 2, item[1]).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.merge_cells("A15:B15"); ws["A15"] = "Como usar este ficheiro"; ws["A15"].fill = PatternFill("solid", fgColor=TEAL); ws["A15"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws.merge_cells("A16:B19")
    ws["A16"] = ("Operacional usa apenas o primeiro snapshot guardado de cada jogo, pelo que repetir uma run ou gerar HTML novo não multiplica o histórico. "
                 "WTA histórico usa a cache local tennis-data.co.uk para estudo factual. As linhas de handicap são referências internas BO3 e não resultados de mercados reais. "
                 "PAPER e REAL devem ser consultados nos seus registos próprios; este ficheiro não prova lucro futuro.")
    ws["A16"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A16"].fill = PatternFill("solid", fgColor="F4F8FC")
    _widths(ws, [40, 68])

    events = wb.create_sheet("Partidas canónicas")
    event_rows = payload["operational"]["event_rows"]
    _write_rows(events, 1, ["ID canónico", "Snapshot", "Analisado UTC", "Início UTC", "Tour", "Torneio", "Jogador A", "Jogador B", "Odd A", "Odd B", "Vencedor", "Score", "Fenzobot favorece"], event_rows,
                ["event_id", "snapshot_key", "analyzed_at_utc", "commence_time_utc", "tour", "tournament", "player_a", "player_b", "odd_a", "odd_b", "winner_side", "result", "fenzobot_side"])
    _widths(events, [42, 28, 22, 22, 10, 34, 25, 25, 10, 10, 10, 20, 25])

    reports = wb.create_sheet("Registo HTML bruto")
    _title(reports, "Registo de relatórios HTML", "Inventário de versões publicadas. Estes ficheiros não são usados nas métricas e podem repetir a mesma partida.")
    _write_rows(reports, 5, ["Data indicada", "Ficheiro HTML"], payload["report_registry"], ["report_date", "report_file"])
    _widths(reports, [18, 100])

    for title, rows, headers, fields, widths, pct_indices in (
        ("Operacional - jogador odd", payload["operational"]["player_odds"], ["Jogador", "Faixa de odd", "Papel", "Jogos", "Vitórias", "Derrotas", "% vitória"], ["player", "odds_band", "role", "matches", "wins", "losses", "win_pct"], [28, 16, 14, 12, 12, 12, 14], [7]),
        ("Operacional - Fenzobot odd", payload["operational"]["fenzobot_odds"], ["Seleção Fenzobot", "Faixa de odd", "Papel", "Jogos", "Vitórias", "Derrotas", "% acerto"], ["player", "odds_band", "role", "matches", "wins", "losses", "win_pct"], [28, 16, 14, 12, 12, 12, 14], [7]),
        ("WTA histórico - jogador odd", payload["historical_wta"]["player_odds"], ["Jogadora", "Faixa de odd", "Papel", "Jogos", "Vitórias", "Derrotas", "% vitória"], ["player", "odds_band", "role", "matches", "wins", "losses", "win_pct"], [28, 16, 14, 12, 12, 12, 14], [7]),
        ("WTA histórico — handicap", payload["historical_wta"]["handicap_reference"], ["Jogadora", "Papel", "Linha referência", "Jogos", "Cobre", "Devolve", "Falha", "% cobre"], ["player", "role", "reference_line", "matches", "covers", "pushes", "fails", "cover_pct"], [28, 14, 18, 12, 12, 12, 12, 14], [8]),
        ("WTA histórico — recuperação", payload["historical_wta"]["set1_recovery"], ["Jogadora", "Perdeu 1.º set", "Recuperou e venceu", "% recuperação"], ["player", "lost_first", "recovered", "recovery_pct"], [28, 18, 22, 16], [4]),
        ("WTA histórico — set decisivo", payload["historical_wta"]["deciding_set"], ["Jogadora", "Sets decisivos", "Venceu", "% vitória"], ["player", "matches", "wins", "win_pct"], [28, 18, 14, 16], [4]),
        ("WTA histórico — tiebreak", payload["historical_wta"]["tiebreak"], ["Jogadora", "Tiebreaks", "Venceu", "% vitória"], ["player", "matches", "wins", "win_pct"], [28, 16, 14, 16], [4]),
    ):
        sheet = wb.create_sheet(title)
        last = _write_rows(sheet, 1, headers, rows, fields)
        _percent_columns(sheet, 2, last, pct_indices)
        _widths(sheet, widths)

    rankings = wb.create_sheet("Rankings")
    _write_rankings(rankings, payload["rankings"])

    raw = wb.create_sheet("WTA histórico bruto")
    raw_rows = []
    for item in raw_wta:
        raw_rows.append({
            "date": item.get("date"), "winner": item.get("winner"), "loser": item.get("loser"),
            "winner_odd": item.get("winner_odd"), "loser_odd": item.get("loser_odd"),
            "winner_games": item.get("winner_games"), "loser_games": item.get("loser_games"),
            "winner_lost_first": item.get("winner_lost_first"), "deciding_set": item.get("deciding_set"),
        })
    _write_rows(raw, 1, ["Data", "Vencedora", "Derrotada", "Odd vencedora", "Odd derrotada", "Games vencedora", "Games derrotada", "Vencedora perdeu 1.º set", "Set decisivo"], raw_rows,
                ["date", "winner", "loser", "winner_odd", "loser_odd", "winner_games", "loser_games", "winner_lost_first", "deciding_set"])
    _widths(raw, [14, 28, 28, 14, 14, 16, 16, 24, 16])

    methodology = wb.create_sheet("Metodologia e limites")
    _title(methodology, "Metodologia e limites", "Leitura correta dos universos e das métricas.")
    rows = [
        ("Deduplicação", "Uma partida entra uma vez: o primeiro snapshot pré-jogo válido. HTMLs repetidos são ficheiros de publicação, não observações analíticas."),
        ("Operacional", "Métricas de Fenzobot por odd usam apenas snapshots canónicos já liquidados. São observacionais, não backtest."),
        ("Histórico WTA", "Resultados e odds vêm apenas das cópias tennis-data.co.uk existentes localmente. Não houve descarga nova nesta construção."),
        ("Handicaps", "Cobertura é calculada contra uma linha interna de referência BO3, inferida pela faixa da Moneyline. Não é uma odd/linha efetivamente oferecida por bookmaker."),
        ("Recuperação", "Conta vitórias após perder o 1.º set no histórico WTA. Não há estatística ponto-a-ponto de breaks."),
        ("Set decisivo e tiebreak", "Calculados com scores completos WTA. Não são disponíveis como universo ATP bruto completo neste checkout."),
        ("PAPER / REAL", "Não são misturados com estes agregados. Usar a carteira PAPER e a folha 22Bet para resultados financeiros."),
    ]
    _header(methodology, 5, ["Tema", "Regra"])
    for index, row in enumerate(rows, 6):
        methodology.cell(index, 1, row[0]).font = Font(name="Arial", bold=True, color="243B53")
        methodology.cell(index, 2, row[1]).alignment = Alignment(wrap_text=True, vertical="top")
        methodology.row_dimensions[index].height = 42
    _widths(methodology, [28, 110])

    # One row per odds band makes the overview readable; the previous chart
    # repeated the same band for every player and was therefore misleading.
    chart_data = wb.create_sheet("Dados gráficos")
    chart_data.sheet_state = "hidden"
    _header(chart_data, 1, ["Faixa", "Acerto Fenzobot", "Decisões liquidadas"])
    chart_rows = [row for row in payload["operational"]["fenzobot_band_summary"] if row["matches"] >= 5]
    for index, row in enumerate(chart_rows, 2):
        chart_data.cell(index, 1, row["odds_band"])
        chart_data.cell(index, 2, (row["win_pct"] or 0) / 100)
        chart_data.cell(index, 3, row["matches"])
    if chart_rows:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Acerto Fenzobot por faixa de odd (n ≥ 5)"
        chart.y_axis.title = "% acerto"
        chart.height = 8
        chart.width = 13
        chart.add_data(Reference(chart_data, min_col=2, min_row=1, max_row=len(chart_rows) + 1), titles_from_data=True)
        chart.set_categories(Reference(chart_data, min_col=1, min_row=2, max_row=len(chart_rows) + 1))
        ws.add_chart(chart, "D5")
    return wb


def _prospective_workbook(baseline: bytes, payload: Mapping[str, Any]) -> Workbook:
    """Append only post-T0 sheets; every historical sheet/cell stays untouched."""
    workbook = load_workbook(BytesIO(baseline), data_only=False)
    for name in ("Pós-T0 - Resumo", "Pós-T0 - Partidas"):
        if name in workbook.sheetnames:
            del workbook[name]
    summary = workbook.create_sheet("Pós-T0 - Resumo")
    _title(
        summary, "Contribuições pós-T0",
        "Universo prospetivo separado; não recalcula nem corrige as folhas históricas.",
    )
    _header(summary, 5, ["Indicador", "Valor"])
    values = payload.get("summary") or {}
    for row, (label, value) in enumerate((
        ("Snapshots pós-T0", values.get("raw_snapshots", 0)),
        ("Partidas canónicas pós-T0", values.get("canonical_snapshots", 0)),
        ("Liquidadas pós-T0", values.get("settled_canonical_snapshots", 0)),
        ("Versões HTML pós-T0", values.get("raw_report_html_versions", 0)),
    ), 6):
        summary.cell(row, 1, label)
        summary.cell(row, 2, value)
    _widths(summary, [42, 24])
    events = workbook.create_sheet("Pós-T0 - Partidas")
    rows = (payload.get("operational") or {}).get("event_rows") or []
    _write_rows(
        events, 1,
        ["ID canónico", "Snapshot", "Analisado UTC", "Início UTC", "Tour", "Torneio", "Jogador A", "Jogador B", "Odd A", "Odd B", "Vencedor", "Score", "Fenzobot favorece"],
        rows,
        ["event_id", "snapshot_key", "analyzed_at_utc", "commence_time_utc", "tour", "tournament", "player_a", "player_b", "odd_a", "odd_b", "winner_side", "result", "fenzobot_side"],
    )
    _widths(events, [42, 28, 22, 22, 10, 34, 25, 25, 10, 10, 10, 20, 25])
    return workbook


def _prospective_history(
    snapshots: Mapping[str, Any], report_files: list[str], boundary: forward_only.Boundary,
) -> dict[str, Any]:
    eligible = []
    for snapshot in snapshots.get("snapshots") or []:
        if isinstance(snapshot, Mapping) and boundary.new_record_eligibility(
            "snapshots", snapshot,
        )[0]:
            eligible.append(dict(snapshot))
    protected = {
        Path(str(item.get("path") or "")).name
        for item in (((boundary.manifest or {}).get("protected") or {}).get(
            "published_reports", {}
        ).get("files") or [])
        if isinstance(item, Mapping)
    }
    future_reports = [name for name in report_files if name not in protected]
    return build_system_history({"snapshots": eligible}, [], future_reports)


def main() -> int:
    snapshots = _read_json(SNAPSHOTS)
    raw_wta = load_local_wta_history(HISTORY_CACHE)
    report_files = [path.name for path in REPORTS_DIR.glob("*.html")]
    payload = build_system_history(snapshots, raw_wta, report_files)
    boundary = forward_only.load_boundary_for_store(ANALYTICS_JSON)
    if boundary.fail_closed:
        raise RuntimeError(boundary.reason_code)
    if boundary.active:
        prospective = _prospective_history(snapshots, report_files, boundary)
        payload = forward_only_projections.compose_json(
            boundary=boundary,
            relative_path=ANALYTICS_JSON.relative_to(ROOT).as_posix(),
            prospective=prospective,
            root=ROOT,
        )
    if boundary.active:
        serialized = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
        try:
            changed = ANALYTICS_JSON.read_text(encoding="utf-8") != serialized
        except OSError:
            changed = True
        if changed:
            ANALYTICS_JSON.parent.mkdir(parents=True, exist_ok=True)
            ANALYTICS_JSON.write_text(serialized, encoding="utf-8")
    else:
        changed = write_if_changed(payload, ANALYTICS_JSON)
    if not changed and WORKBOOK_PATH.exists():
        print("[info] histórico sem alterações semânticas; Excel preservado.")
        return 0
    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    if boundary.active:
        baseline = forward_only_projections.baseline_bytes(
            boundary, WORKBOOK_PATH.relative_to(ROOT).as_posix(), root=ROOT,
        )
        wb = _prospective_workbook(baseline, prospective)
    else:
        wb = _workbook(payload, raw_wta)
    wb.save(WORKBOOK_PATH)
    print(f"[info] histórico canónico atualizado: {WORKBOOK_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
