from __future__ import annotations

import hashlib
import json
import os
import subprocess
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from scripts.build_system_history_workbook import _prospective_workbook
from src import (
    dashboard,
    forward_only,
    forward_only_projections,
    green_monetization,
    green_strong_validation,
    main,
    market_memory_report,
)


def git(root: Path, *args: str) -> str:
    return subprocess.check_output(["git", *args], cwd=root, text=True).strip()


def active_manifest(
    path: Path, *, commit: str, references: list[dict], mutable_indexes: list[dict] | None = None,
) -> None:
    empty = {"count": 0, "records": [], "identity_tokens": [], "weak_alias_contexts": {}}
    document = {
        "schema_version": 1, "change_id": forward_only.CHANGE_ID, "status": "ACTIVE",
        "effective_from_utc": "2026-09-30T10:00:00+00:00",
        "code_commit": commit, "data_base_commit": commit,
        "protected": {
            "snapshots": dict(empty), "paper": dict(empty), "market_ledger": dict(empty),
            "published_reports": {"count": 0, "files": []}, "exclusions": [],
            "historic_aggregate_references": references,
            "mutable_index_versions": mutable_indexes or [],
        },
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(document), encoding="utf-8")
    forward_only.activation_lock_path(path).write_text(json.dumps({
        "schema_version": 1, "change_id": forward_only.CHANGE_ID,
        "effective_from_utc": document["effective_from_utc"],
        "manifest_sha256": forward_only.canonical_sha256(document),
        "code_commit": commit, "data_base_commit": commit,
    }), encoding="utf-8")


class ForwardOnlyProjectionTests(unittest.TestCase):
    def test_imprecise_json_and_xlsx_baseline_stay_exact_with_future_component(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            git(root, "init", "-b", "main")
            git(root, "config", "user.name", "test")
            git(root, "config", "user.email", "test@example.invalid")
            json_path = root / "data/dashboard/system_history_analytics.json"
            json_path.parent.mkdir(parents=True)
            baseline_json = {
                "summary": {"canonical_snapshots": 999, "settled_canonical_snapshots": 123},
                "input_fingerprint_sha256": "intentionally-imprecise",
            }
            json_path.write_text(json.dumps(baseline_json), encoding="utf-8")
            xlsx_path = root / "docs/dashboard/Fenzobot_Historico_do_Sistema.xlsx"
            xlsx_path.parent.mkdir(parents=True)
            workbook = Workbook(); workbook.active.title = "Resumo"
            workbook["Resumo"]["A1"] = "baseline-imprecisa"
            workbook["Resumo"]["B2"] = "=1+1"
            workbook.save(xlsx_path)
            git(root, "add", "."); git(root, "commit", "-m", "baseline")
            commit = git(root, "rev-parse", "HEAD")
            references = [
                {"path": json_path.relative_to(root).as_posix(), "kind": "JSON", "sha256_at_t0": hashlib.sha256(json_path.read_bytes()).hexdigest()},
                {"path": xlsx_path.relative_to(root).as_posix(), "kind": "XLSX", "sha256_at_t0": hashlib.sha256(xlsx_path.read_bytes()).hexdigest()},
            ]
            manifest = root / "manifest.json"
            active_manifest(manifest, commit=commit, references=references)
            boundary = forward_only.load_boundary(manifest)
            zero = forward_only_projections.compose_json(
                boundary=boundary, relative_path=references[0]["path"],
                prospective={"summary": {"canonical_snapshots": 0}}, root=root,
            )
            self.assertEqual(
                forward_only_projections.historic_component(zero), baseline_json,
            )
            one = forward_only_projections.compose_json(
                boundary=boundary, relative_path=references[0]["path"],
                prospective={"summary": {"canonical_snapshots": 1}}, root=root,
            )
            self.assertEqual(forward_only_projections.historic_component(one), baseline_json)
            self.assertEqual(
                one["forward_only_continuity"]["prospective"]["summary"]["canonical_snapshots"], 1,
            )
            rebuilt = _prospective_workbook(
                forward_only_projections.baseline_bytes(boundary, references[1]["path"], root=root),
                {"summary": {}, "operational": {"event_rows": []}},
            )
            output = BytesIO(); rebuilt.save(output)
            reread = load_workbook(BytesIO(output.getvalue()), data_only=False)
            self.assertEqual(reread["Resumo"]["A1"].value, "baseline-imprecisa")
            self.assertEqual(reread["Resumo"]["B2"].value, "=1+1")
            self.assertIn("Pós-T0 - Resumo", reread.sheetnames)

    def test_midday_index_is_versioned_and_future_index_remains_writable(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            git(root, "init", "-b", "main"); git(root, "config", "user.name", "test"); git(root, "config", "user.email", "test@example.invalid")
            reports = root / "docs/relatorios"; frozen = reports / "frozen-indexes"
            frozen.mkdir(parents=True)
            frozen_file = frozen / "index-2026-09-30-t0-abc.html"
            frozen_file.write_bytes(b"historic-index-bytes")
            git(root, "add", "."); git(root, "commit", "--allow-empty", "-m", "base")
            commit = git(root, "rev-parse", "HEAD")
            manifest = root / "manifest.json"
            active_manifest(manifest, commit=commit, references=[], mutable_indexes=[{
                "source_path": "docs/relatorios/index-2026-09-30.html",
                "path": "docs/relatorios/frozen-indexes/index-2026-09-30-t0-abc.html",
                "sha256": hashlib.sha256(frozen_file.read_bytes()).hexdigest(),
            }])
            with patch.dict(os.environ, {"FENZOBOT_FORWARD_ONLY_MANIFEST": str(manifest)}), patch.object(main, "SITE_OUTPUT_DIR", str(root / "docs")):
                main._write_site_index([], "2026-09-30", str(reports))
            self.assertEqual(frozen_file.read_bytes(), b"historic-index-bytes")
            current = (reports / "index-2026-09-30.html").read_text(encoding="utf-8")
            self.assertIn("Versão histórica congelada em T0", current)
            self.assertIn("frozen-indexes/index-2026-09-30-t0-abc.html", current)

    def test_real_projection_writers_preserve_intentionally_imprecise_baselines(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            git(root, "init", "-b", "main"); git(root, "config", "user.name", "test"); git(root, "config", "user.email", "test@example.invalid")
            paths = {
                "memory": root / "data/market_ledger/derived/market-memory-v1.json",
                "green": root / "data/validation/green-strong-v1.json",
                "money": root / "data/validation/green-monetization-v1.json",
                "dashboard": root / "data/dashboard/fenzobot-dashboard-v1.json",
            }
            baselines = {
                "memory": {"observation_count": 777, "evaluation": {"legacy_rate": 12.34}},
                "green": {"sample": {"candidates": 888}, "legacy_rate": 43.21},
                "money": {"eligible_entries": 999, "roi_pct": -91.0},
                "dashboard": {
                    "semantic_fingerprint": "baseline-fingerprint",
                    "generated_at_utc": "2026-09-30T09:00:00+00:00",
                    "green_monetization_v1": {"generated_at_utc": "2026-09-30T09:00:00+00:00"},
                    "global": {"total_reports": 456},
                },
            }
            for name, path in paths.items():
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(json.dumps(baselines[name]), encoding="utf-8")
            (root / "data/calibration_snapshots.json").write_text('{"snapshots":[]}\n', encoding="utf-8")
            (root / "data/paper_trades.json").write_text('{"entries":[]}\n', encoding="utf-8")
            (root / "data/paper_integrity_exclusions.json").write_text('{"exclusions":[]}\n', encoding="utf-8")
            (root / "data/validation/market-integrity-exclusions-v1.json").write_text('{"exclusions":[]}\n', encoding="utf-8")
            (root / "docs/relatorios").mkdir(parents=True)
            git(root, "add", "."); git(root, "commit", "-m", "imprecise baseline")
            commit = git(root, "rev-parse", "HEAD")
            references = [{
                "path": path.relative_to(root).as_posix(), "kind": "JSON",
                "sha256_at_t0": hashlib.sha256(path.read_bytes()).hexdigest(),
            } for path in paths.values()]
            manifest = root / "manifest.json"
            active_manifest(manifest, commit=commit, references=references)
            memory = market_memory_report.build_and_write(
                ledger_root=root / "data/market_ledger",
                snapshots_path=root / "data/calibration_snapshots.json",
                paper_path=root / "data/paper_trades.json",
                output_path=paths["memory"], protection_manifest_path=manifest,
                continuity_root=root,
            )
            green = green_strong_validation.build_and_write(
                memory_report=memory, manual_path=root / "missing-manual.json",
                output_path=paths["green"], protection_manifest_path=manifest,
                continuity_root=root,
            )
            money = green_monetization.build_and_write(
                paper_path=root / "data/paper_trades.json",
                legacy_exclusions_path=root / "data/paper_integrity_exclusions.json",
                market_exclusions_path=root / "data/validation/market-integrity-exclusions-v1.json",
                output_path=paths["money"], protection_manifest_path=manifest,
                continuity_root=root,
            )
            control = dashboard.build_and_write(
                root=root, protection_manifest_path=manifest,
                output_path=paths["dashboard"], html_path=root / "docs/dashboard/index.html",
            )
            for name, document in (
                ("memory", memory), ("green", green), ("money", money),
                ("dashboard", control),
            ):
                self.assertEqual(
                    forward_only_projections.historic_component(document), baselines[name],
                )
                self.assertIn("prospective", document["forward_only_continuity"])

    def test_google_sync_routes_active_payload_only_to_post_t0_sheets(self):
        source = (Path(__file__).parents[1] / "scripts/google_apps_script/sync_system_history.gs").read_text(encoding="utf-8")
        self.assertIn("writeProspectiveSystemHistory_", source)
        function = source[source.index("function writeProspectiveSystemHistory_"):source.index("function installSystemHistorySync")]
        self.assertNotIn("writeSystemHistoryWorkbook_(", function)
        self.assertIn("Pós-T0 - Resumo", function)
        self.assertIn("Pós-T0 - Partidas", function)


if __name__ == "__main__":
    unittest.main()
