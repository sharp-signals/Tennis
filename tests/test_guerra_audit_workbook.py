import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from scripts.build_system_history_workbook import (
    _stamp_workbook_fingerprint,
    _stored_workbook_fingerprint,
    _write_factor_coverage_sheets,
    _write_guerra_audit_sheets,
)


class GuerraAuditWorkbookTests(unittest.TestCase):
    def test_workbook_fingerprint_detects_stale_or_interrupted_projection(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "history.xlsx"
            workbook = Workbook()
            self.assertEqual(_stored_workbook_fingerprint(path), "")
            _stamp_workbook_fingerprint(workbook, "fingerprint-1")
            workbook.save(path)
            self.assertEqual(_stored_workbook_fingerprint(path), "fingerprint-1")

    def test_writes_factor_coverage_and_reason_tabs(self):
        workbook = Workbook()
        _write_factor_coverage_sheets(workbook, {
            "summary": [{
                "scope": "CHALLENGER", "factor": "servico_recente",
                "seen_snapshots": 10, "available_snapshots": 4,
                "unavailable_snapshots": 6, "available_inactive_snapshots": 1,
                "active_snapshots": 3, "availability_pct": 40.0,
                "active_pct": 30.0,
            }],
            "reasons": [{
                "scope": "CHALLENGER", "factor": "servico_recente",
                "state": "DADO_INDISPONIVEL",
                "reason": "amostra em falta num dos lados", "snapshots": 6,
                "pct_of_factor_snapshots": 60.0,
            }],
        })
        self.assertIn("Cobertura — fatores", workbook.sheetnames)
        self.assertIn("Cobertura — motivos", workbook.sheetnames)
        self.assertEqual(workbook["Cobertura — fatores"]["D6"].value, 4)
        self.assertEqual(workbook["Cobertura — fatores"]["H6"].value, 0.4)
        self.assertEqual(
            workbook["Cobertura — motivos"]["D6"].value,
            "amostra em falta num dos lados",
        )

    def test_writes_aggregate_tabs_without_individual_rows(self):
        workbook = Workbook()
        payload = {
            "by_strategy": {
                "GUERRA_SELECTION_V1": {
                    "audit_v1": {
                        "status": "AVAILABLE",
                        "privacy_contract": "AGGREGATES_ONLY_NO_INDIVIDUAL_SELECTIONS",
                        "quality": {"linked_entries": 12, "canonical_entries": 8},
                        "factors": [{
                            "scope": "ATP", "factor": "forma_recente",
                            "relation": "Alinha com seleção", "total_entries": 8,
                            "settled": 6, "wins": 4, "losses": 2,
                            "win_rate_pct": 66.7, "roi_pct": 12.5,
                            "sample_status": "Informativo",
                        }],
                        "segments": [{
                            "dimension": "Mercado", "group": "Moneyline",
                            "total_entries": 5, "settled": 4, "wins": 3,
                            "losses": 1, "win_rate_pct": 75.0,
                            "sample_status": "Informativo",
                        }],
                    },
                },
            },
        }
        _write_guerra_audit_sheets(workbook, payload)
        self.assertIn("GUERRA — fatores", workbook.sheetnames)
        self.assertIn("GUERRA — segmentos", workbook.sheetnames)
        self.assertIn("GUERRA — qualidade", workbook.sheetnames)
        self.assertEqual(workbook["GUERRA — fatores"]["A6"].value, "ATP")
        self.assertEqual(workbook["GUERRA — fatores"]["I6"].value, 0.667)
        self.assertEqual(workbook["GUERRA — qualidade"]["B8"].value, 12)


if __name__ == "__main__":
    unittest.main()
